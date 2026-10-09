"""Parse university Excel schedule workbook into application JSON format."""

import argparse
import json
import re
from pathlib import Path
from typing import Any, Dict, List, Tuple

import openpyxl

DEFAULT_EXCEL_PATH = Path(
    r"C:\Users\Ak417ytrfg\.t3\userdata\attachments\75dfb82c-8301-414a-96c7-78c6eb493195-003e016b-5626-41b8-931c-860c98167df2-xlsx.xlsx"
)
DEFAULT_OUTPUT_PATH = Path("src/data/schedule.json")

GROUP_REGEX = re.compile(r"^[А-ЯA-Z]{2,}-[А-ЯA-Z]-\d{2}/\d")

DAY_NAMES_MAP = {
    "ПОНЕДЕЛЬНИК": "Понедельник",
    "ВТОРНИК": "Вторник",
    "СРЕДА": "Среда",
    "ЧЕТВЕРГ": "Четверг",
    "ПЯТНИЦА": "Пятница",
    "СУББОТА": "Суббота",
}

DAYS_ORDER = [
    "Понедельник",
    "Вторник",
    "Среда",
    "Четверг",
    "Пятница",
    "Суббота",
]

LESSON_TIMES: Dict[int, str] = {
    1: "08:00 \u2013 09:20",
    2: "09:30 \u2013 10:50",
    3: "11:00 \u2013 12:20",
    4: "12:45 \u2013 14:05",
    5: "14:15 \u2013 15:35",
    6: "15:45 \u2013 17:05",
    7: "17:15 \u2013 18:35",
    8: "18:45 \u2013 20:05",
}


def clean_text(raw_value: Any) -> str:
    """Normalize text value and remove extra whitespace."""
    if raw_value is None:
        return ""
    text = str(raw_value).replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text or text.lower() == "nan":
        return ""
    return text


def detect_day_and_lesson_columns(sheet: Any) -> Tuple[int, int]:
    """Locate column indices for days and lesson numbers."""
    day_col = None
    for r in range(1, min(15, sheet.max_row + 1)):
        for c in range(1, sheet.max_column + 1):
            val = sheet.cell(r, c).value
            if val and str(val).strip().upper() in DAY_NAMES_MAP:
                day_col = c
                break
        if day_col:
            break

    if not day_col:
        raise ValueError(f"Day column not detected in sheet {sheet.title}")

    # Check candidates for lesson column
    lesson_col = None
    candidates = [day_col + 1, day_col - 1, day_col]
    for cand in candidates:
        if 1 <= cand <= sheet.max_column:
            found_nums = set()
            for r in range(1, sheet.max_row + 1):
                val = sheet.cell(r, cand).value
                if val is not None and str(val).strip().isdigit():
                    num = int(str(val).strip())
                    if 1 <= num <= 8:
                        found_nums.add(num)
            if len(found_nums) >= 5:
                lesson_col = cand
                break

    if not lesson_col:
        lesson_col = day_col + 1

    return day_col, lesson_col


def detect_group_columns(sheet: Any) -> Dict[int, str]:
    """Find group headers in rows 1 to 8."""
    groups_by_col: Dict[int, str] = {}
    for r in range(1, min(9, sheet.max_row + 1)):
        for c in range(1, sheet.max_column + 1):
            val = sheet.cell(r, c).value
            if val is not None:
                cleaned = str(val).strip()
                if GROUP_REGEX.match(cleaned):
                    groups_by_col[c] = cleaned
    return groups_by_col


def resolve_course(sheet_name: str, group_name: str) -> str:
    """Determine course identifier from sheet name and group name."""
    if sheet_name.startswith("ФИТТ-1"):
        return "1"
    if sheet_name.startswith("ФИТТ-2"):
        return "2"
    if sheet_name.startswith("ФИТТ-3"):
        return "3"
    if sheet_name.startswith("ФИТТ-4"):
        return "4"
    if sheet_name.startswith("ФИТТ-5"):
        return "5"
    if sheet_name.startswith("ФИТТ-М"):
        if "26/" in group_name:
            return "Магистратура 1"
        if "25/" in group_name:
            return "Магистратура 2"
        raise ValueError(
            f"Cannot determine magistracy course for group: {group_name}"
        )
    raise ValueError(f"Unknown course pattern for sheet: {sheet_name}")


def parse_schedule(excel_path: Path) -> Dict[str, Any]:
    """Parse Excel workbook into structured schedule dictionary."""
    wb = openpyxl.load_workbook(str(excel_path), data_only=True)
    raw_schedule: Dict[str, Dict[str, Dict[str, Dict[str, Dict[str, List[Dict[str, Any]]]]]]] = {}

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        week = "1" if sheet_name.endswith("1") else "2"
        day_col, lesson_col = detect_day_and_lesson_columns(sheet)
        groups = detect_group_columns(sheet)

        if not groups:
            continue

        # Map sheet rows to day and lesson number
        curr_day = None
        row_mapping: Dict[int, Tuple[str, int]] = {}
        for r in range(1, sheet.max_row + 1):
            day_cell_val = sheet.cell(r, day_col).value
            if day_cell_val is not None:
                day_str = str(day_cell_val).strip().upper()
                if day_str in DAY_NAMES_MAP:
                    curr_day = DAY_NAMES_MAP[day_str]

            lesson_cell_val = sheet.cell(r, lesson_col).value
            if curr_day and lesson_cell_val is not None:
                lesson_str = str(lesson_cell_val).strip()
                if lesson_str.isdigit():
                    lesson_num = int(lesson_str)
                    if 1 <= lesson_num <= 8:
                        row_mapping[r] = (curr_day, lesson_num)

        # Parse lessons for each group in this sheet
        for col_idx, group_name in groups.items():
            category = group_name.split("-")[0]
            course = resolve_course(sheet_name, group_name)

            group_data = (
                raw_schedule
                .setdefault(category, {})
                .setdefault(course, {})
                .setdefault(group_name, {"1": {}, "2": {}})
            )

            for r, (day, lesson_num) in row_mapping.items():
                cell_value = sheet.cell(r, col_idx).value
                cleaned_lesson_text = clean_text(cell_value)
                if not cleaned_lesson_text:
                    continue

                lesson_entry = {
                    "num": lesson_num,
                    "time": LESSON_TIMES[lesson_num],
                    "text": cleaned_lesson_text,
                }
                group_data[week].setdefault(day, []).append(lesson_entry)

    # Sort dictionary entries deterministically
    sorted_schedule: Dict[str, Any] = {}
    for category in sorted(raw_schedule.keys()):
        sorted_schedule[category] = {}
        courses = raw_schedule[category]

        def course_key(crs: str) -> Tuple[int, Any]:
            if crs.isdigit():
                return (0, int(crs))
            return (1, crs)

        for course in sorted(courses.keys(), key=course_key):
            sorted_schedule[category][course] = {}
            groups = courses[course]
            for group_name in sorted(groups.keys()):
                group_data = groups[group_name]
                sorted_schedule[category][course][group_name] = {"1": {}, "2": {}}
                for wk in ["1", "2"]:
                    days_dict = group_data.get(wk, {})
                    sorted_days: Dict[str, List[Dict[str, Any]]] = {}
                    for d in DAYS_ORDER:
                        if d in days_dict and days_dict[d]:
                            sorted_lessons = sorted(days_dict[d], key=lambda x: x["num"])
                            sorted_days[d] = sorted_lessons
                    sorted_schedule[category][course][group_name][wk] = sorted_days

    return sorted_schedule


def count_statistics(data: Dict[str, Any]) -> Dict[str, int]:
    """Compute aggregate counts from schedule data."""
    total_categories = len(data)
    total_courses = sum(len(courses) for courses in data.values())
    total_groups = sum(
        sum(len(groups) for groups in courses.values())
        for courses in data.values()
    )
    total_lessons = 0
    for courses in data.values():
        for groups in courses.values():
            for weeks in groups.values():
                for days in weeks.values():
                    for lessons in days.values():
                        total_lessons += len(lessons)

    return {
        "categories": total_categories,
        "courses": total_courses,
        "groups": total_groups,
        "lessons": total_lessons,
    }


def main() -> None:
    """Execute Excel to JSON schedule conversion."""
    parser = argparse.ArgumentParser(
        description="Convert university Excel schedule to JSON."
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=DEFAULT_EXCEL_PATH,
        help="Path to input Excel schedule file.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT_PATH,
        help="Path to output schedule JSON file.",
    )
    args = parser.parse_args()

    input_path = args.input.resolve()
    output_path = args.output.resolve()

    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    print(f"Reading schedule from: {input_path}")
    data = parse_schedule(input_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f"Schedule successfully written to: {output_path}")

    stats = count_statistics(data)
    print("\nParsing Statistics:")
    print(f"- Total categories: {stats['categories']}")
    print(f"- Total courses: {stats['courses']}")
    print(f"- Total groups: {stats['groups']}")
    print(f"- Total lessons: {stats['lessons']}")


if __name__ == "__main__":
    main()
